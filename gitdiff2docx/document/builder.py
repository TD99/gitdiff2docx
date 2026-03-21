"""Document builder utilities."""

from datetime import datetime
from docx import Document
from gitdiff2docx.document.tables import add_legend_table


def create_document(commit1, commit2, lang, config, theme):
    """Create a new document with initial headers and legend."""
    doc = Document()
    doc.add_heading(f"{lang['git_changes_report']} ({commit1} → {commit2})", level=1)
    doc.add_paragraph(lang["report_generated_on"].format(
        date=datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ))

    # Add legend table
    doc.add_heading(lang["legend"], level=config.get("heading_level", 2))
    add_legend_table(doc, theme, config, lang)

    doc.add_page_break()
    doc.add_heading(lang["diffs"], level=config.get("heading_level", 2))
    
    return doc
