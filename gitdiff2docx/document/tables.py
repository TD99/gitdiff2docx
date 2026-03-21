"""Document table generation for diffs and legends."""

from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from pygments import lex

from gitdiff2docx.utils.color import rgb_from_hex
from gitdiff2docx.document.layout import get_usable_width, apply_table_theme_style
from gitdiff2docx.theme.styling import get_theme_font


def add_legend_table(document, theme, config, lang):
    """Add a legend table to the document."""
    legend_table = document.add_table(rows=0, cols=2)
    legend_table.style = "Table Grid"
    apply_table_theme_style(legend_table, theme)

    theme_colors = theme.get("colors", {})
    theme_symbols = theme.get("symbols", {})
    theme_font = theme.get("font", {})

    legend_add_color = theme_colors.get("add_fill", "D0FFD0")
    legend_remove_color = theme_colors.get("remove_fill", "FFD0D0")
    legend_neutral_color = theme_colors.get("neutral_fill", "F5F5F5")

    legend_add_symbol = theme_symbols.get("add", "+")
    legend_remove_symbol = theme_symbols.get("remove", "-")
    legend_neutral_symbol = theme_symbols.get("neutral", "=")

    diff_font_name, diff_font_size = get_theme_font(theme, config)
    bold_symbols = bool(theme_font.get("bold_symbols", False))

    legend_data = [
        (lang["legend_add"], legend_add_color, legend_add_symbol),
        (lang["legend_remove"], legend_remove_color, legend_remove_symbol),
    ]

    # Only include neutral/unchanged lines in the legend if they're being shown in the diff
    if config.get("include_unchanged_lines", True):
        legend_data.append(
            (lang["legend_neutral"], legend_neutral_color, legend_neutral_symbol)
        )

    for label, color, symbol in legend_data:
        column = legend_table.add_row().cells
        column[0].text = label

        shading = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls("w"), color))
        column[1]._element.get_or_add_tcPr().append(shading)
        column[1].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        p = column[1].paragraphs[0]
        run = p.add_run(symbol)
        font = run.font
        font.name = diff_font_name
        font.size = Pt(diff_font_size)
        if bold_symbols:
            run.bold = True
        if symbol == legend_add_symbol:
            run.font.color.rgb = rgb_from_hex(theme_colors.get("add_symbol", "000000"))
        elif symbol == legend_remove_symbol:
            run.font.color.rgb = rgb_from_hex(theme_colors.get("remove_symbol", "000000"))
        else:
            run.font.color.rgb = rgb_from_hex(theme_colors.get("neutral_symbol", "000000"))


def add_diff_table(document, diff_lines, lexer, theme, config, token_styles):
    """Add a formatted and syntax-highlighted code diff table."""
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    apply_table_theme_style(table, theme)

    theme_colors = theme.get("colors", {})
    theme_symbols = theme.get("symbols", {})
    theme_font = theme.get("font", {})

    symbol_col_width = Cm(float(theme.get("symbol_column_width_cm", 0.57)))
    table.columns[0].width = symbol_col_width
    table.columns[1].width = get_usable_width(document) - symbol_col_width

    diff_font_name, diff_font_size = get_theme_font(theme, config)
    diff_font_size_pt = Pt(diff_font_size)

    add_symbol = theme_symbols.get("add", "+")
    remove_symbol = theme_symbols.get("remove", "-")
    neutral_symbol = theme_symbols.get("neutral", "=")

    bold_symbols = bool(theme_font.get("bold_symbols", False))
    center_symbols = bool(theme_font.get("center_symbols", False))
    line_spacing = float(theme_font.get("line_spacing", 1.0))
    use_syntax_highlighting = bool(theme.get("use_syntax_highlighting", True))

    # Skip unchanged lines if configured to do so
    include_unchanged = config.get("include_unchanged_lines", True)
    if not include_unchanged:
        diff_lines = [line for line in diff_lines if not line.startswith(" ")]

    for line in diff_lines:
        row_cells = table.add_row().cells
        symbol_cell = row_cells[0]
        code_cell = row_cells[1]

        if line.startswith("+"):
            fill = theme_colors.get("add_fill", "D0FFD0")
            symbol_color = theme_colors.get("add_symbol", "000000")
            symbol = add_symbol
        elif line.startswith("-"):
            fill = theme_colors.get("remove_fill", "FFD0D0")
            symbol_color = theme_colors.get("remove_symbol", "000000")
            symbol = remove_symbol
        else:
            fill = theme_colors.get("neutral_fill", "F5F5F5")
            symbol_color = theme_colors.get("neutral_symbol", "000000")
            symbol = neutral_symbol

        for cell in (symbol_cell, code_cell):
            fill_normalized = fill.lstrip("#") if isinstance(fill, str) else fill
            shading = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls("w"), fill_normalized))
            cell._element.get_or_add_tcPr().append(shading)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        symbol_paragraph = symbol_cell.paragraphs[0]
        symbol_paragraph.clear()
        symbol_paragraph.paragraph_format.space_before = Pt(0)
        symbol_paragraph.paragraph_format.space_after = Pt(0)
        if center_symbols:
            symbol_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_sym = symbol_paragraph.add_run(symbol)
        run_sym.font.name = diff_font_name
        run_sym.font.size = diff_font_size_pt
        run_sym.bold = bold_symbols
        run_sym.font.color.rgb = rgb_from_hex(symbol_color)

        paragraph = code_cell.paragraphs[0]
        paragraph.clear()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        if line_spacing > 0:
            paragraph.paragraph_format.line_spacing = line_spacing

        code_content = line[1:]

        for ttype, value in lex(code_content, lexer):
            value = value.rstrip('\n')
            if not value:
                value = "\u00A0"

            run = paragraph.add_run(value)
            run.font.name = diff_font_name
            run.font.size = diff_font_size_pt

            style_str = token_styles.get(ttype) if use_syntax_highlighting else None
            if style_str:
                for part in style_str.split():
                    if part == "bold":
                        run.bold = True
                    elif part == "italic":
                        run.italic = True
                    elif part.startswith("#") and len(part) == 7:
                        hexcode = part.lstrip("#")
                        try:
                            r = int(hexcode[0:2], 16)
                            g = int(hexcode[2:4], 16)
                            b = int(hexcode[4:6], 16)
                            run.font.color.rgb = RGBColor(r, g, b)
                        except ValueError:
                            pass
            else:
                run.font.color.rgb = rgb_from_hex(theme_colors.get("default_code", "000000"))
